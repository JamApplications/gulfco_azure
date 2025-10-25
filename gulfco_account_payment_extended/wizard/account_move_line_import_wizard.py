import base64
import io
import openpyxl
from odoo import models, fields
from odoo.exceptions import UserError

class AccountMoveLineImportWizard(models.TransientModel):
    _name = 'account.move.line.import.wizard'
    _description = 'Import Account Move Line Data'

    upload_file = fields.Binary(string="Upload Excel File", required=True)
    filename = fields.Char("Filename")
    comment = fields.Char(string="Comment")

    def action_import_move_line_data(self):
        if not self.upload_file:
            raise UserError("Please upload a file.")

        # Decode base64 file content
        try:
            file_data = base64.b64decode(self.upload_file)
            file = io.BytesIO(file_data)
            workbook = openpyxl.load_workbook(file, data_only=True)
            sheet = workbook.active
        except Exception as e:
            raise UserError(f"Unable to read Excel file: {e}")

        # Read headers
        headers = [str(cell.value).strip() if cell.value else '' for cell in
                   next(sheet.iter_rows(min_row=1, max_row=1))]
        required_fields = ['id', 'payment_amount', 'matching_number', 'matching_reference']

        if not all(field in headers for field in required_fields):
            raise UserError(f"Excel file must contain the following columns: {', '.join(required_fields)}")

        header_index = {key: headers.index(key) for key in required_fields + ['comment']}
        updated = 0

        move_lines = self.env['account.move.line'].sudo()

        for row in sheet.iter_rows(min_row=2, values_only=True):
            try:
                line_id = int(row[header_index['id']])
                payment_amount = float(row[header_index['payment_amount']] or 0.0)
                matching_number = str(row[header_index['matching_number']] or '').strip()
                matching_reference = str(row[header_index['matching_reference']] or '').strip()
                comment = str(row[header_index['comment']] or '').strip()


                self.env.cr.execute("""
                    UPDATE account_move_line
                    SET payment_amount = %s,
                        matching_number = %s,
                        matching_reference = %s,
                        comment = %s
                    WHERE id = %s
                """, (payment_amount, matching_number, matching_reference,comment, line_id))
                updated += 1

                move_line = self.env['account.move.line'].sudo().browse(line_id)
                if move_line:
                    move_line._compute_payment_amount_currency()

                move_lines = move_lines + move_line



            except Exception as e:
                raise UserError(f"Error on row with ID {row[0]}: {e}")

        for matching_reference in tuple(set(move_lines.mapped('matching_reference'))):
            rec_move_lines = move_lines.filtered(lambda x:x.matching_reference == matching_reference)
            wizard = self.env['account.reconcile.wizard'].with_context(
                active_model='account.move.line',
                active_ids=rec_move_lines.ids,
            ).new({})

            wizard.write({'allow_partials':True})
            wizard.reconcile()
            # if self.comment:
            #     rec_move_lines.sudo().write({'comment':self.comment})

        return {
            'type': 'ir.actions.client',
            'tag': 'soft_reload',
            'params': {
                'title': 'Move Lines Imported',
                'message': f'{updated} move lines were updated and reconciled successfully.',
                'type': 'success',
                'sticky': False,
            }
        }
